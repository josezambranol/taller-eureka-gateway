import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("INFORME DE ENTREGABLE DE TALLER\nIMPLEMENTACIÓN SPRING BOOT + EUREKA SERVER & CLIENT + API GATEWAY")
    run_title.bold = True
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = RGBColor(26, 54, 93) # Navy Blue
    run_title.font.name = "Calibri"

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Dynamic Routing & Service Discovery en Microservicios con Spring Cloud\nBasado en la Guía de GeeksforGeeks")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(74, 85, 104)

    doc.add_paragraph() # Spacing

    # Metadata Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    metadata = [
        ("Autor / Estudiante:", "JOSE DANIEL ZAMBRANO LUNA"),
        ("Proyecto / Entregable:", "Taller Spring Boot + Eureka Server & Client + API Gateway"),
        ("Referencia Técnica:", "https://www.geeksforgeeks.org/advance-java/dynamic-routing-and-service-discovery-in-api-gateway/"),
        ("Tecnologías y Versiones:", "Java 21 LTS, Spring Boot 3.3.5, Spring Cloud 2023.0.3, Maven 3.9"),
        ("Estado de Validación:", "Completado y Verificado 100% en Ejecución Real")
    ]

    for i, (k, v) in enumerate(metadata):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "EDF2F7")
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(10)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)

    doc.add_paragraph()
    doc.add_heading("1. Introducción y Objetivos del Taller", level=1)
    
    p = doc.add_paragraph(
        "En una arquitectura de microservicios distribuida, la gestión de ubicaciones de red (direcciones IP y puertos) "
        "de forma estática resulta inviable a escala debido al escalado dinámico, despliegues continuos y posibles fallos. "
        "Este taller implementa un ecosistema desacoplado y resiliente compuesto por:"
    )
    doc.add_paragraph("• Eureka Server: Servidor centralizado de registro y descubrimiento de servicios.", style='List Bullet')
    doc.add_paragraph("• UserService (Eureka Client): Microservicio proveedor que expone endpoints REST y se auto-registra.", style='List Bullet')
    doc.add_paragraph("• API Gateway: Punto de entrada unificado que implementa enrutamiento dinámico mediante balanceador de carga (lb://USER-SERVICE).", style='List Bullet')

    doc.add_heading("2. Arquitectura de la Solución", level=1)
    p_arch = doc.add_paragraph(
        "El ecosistema opera bajo el patrón de Service Discovery del lado del cliente / gateway:\n"
        "1. Al iniciar, UserService se registra en Eureka Server (http://localhost:9099/eureka) con el nombre 'user-service'.\n"
        "2. ApiGateway se inicia en el puerto 9056, se registra y descarga periódicamente la lista de instancias disponibles.\n"
        "3. Cuando un cliente externo realiza una solicitud GET a http://localhost:9056/client, ApiGateway intercepta la petición.\n"
        "4. Mediante el prefijo 'lb://USER-SERVICE', el Gateway consulta a Eureka la ubicación de las instancias activas y redirige el tráfico transparentemente.\n"
        "5. El cliente recibe la respuesta 'Welcome to client' sin haber tenido conocimiento previo de la IP o puerto del servicio final (8086)."
    )

    doc.add_heading("3. Componentes Implementados y Configuración", level=1)

    doc.add_heading("3.1 Módulo 1: EurekaServerService (Puerto 9099)", level=2)
    doc.add_paragraph(
        "Actúa como el registro de nombres central. Está habilitado mediante la anotación @EnableEurekaServer en la clase principal."
    )
    p_code1 = doc.add_paragraph()
    r = p_code1.add_run("Configuración application.properties:\n"
                      "spring.application.name=EurekaServerService\n"
                      "server.port=9099\n"
                      "eureka.instance.hostname=localhost\n"
                      "eureka.client.register-with-eureka=false\n"
                      "eureka.client.fetch-registry=false\n"
                      "eureka.client.service-url.defaultZone=http://${eureka.instance.hostname}:${server.port}/eureka")
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

    doc.add_heading("3.2 Módulo 2: UserService (Puerto 8086)", level=2)
    doc.add_paragraph(
        "Microservicio de negocio habilitado como cliente con @EnableDiscoveryClient. Implementa el controlador REST UserController."
    )
    p_code2 = doc.add_paragraph()
    r = p_code2.add_run("Configuración application.properties:\n"
                      "spring.application.name=user-service\n"
                      "server.port=8086\n"
                      "eureka.instance.prefer-ip-address=true\n"
                      "eureka.client.fetch-registry=true\n"
                      "eureka.client.register-with-eureka=true\n"
                      "eureka.client.service-url.defaultZone=http://localhost:9099/eureka")
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

    p_code2b = doc.add_paragraph()
    r = p_code2b.add_run("Controlador UserController.java:\n"
                       "@RestController\n"
                       "public class UserController {\n"
                       "    @GetMapping(\"/client\")\n"
                       "    public String check() {\n"
                       "        return \"Welcome to client\";\n"
                       "    }\n"
                       "}")
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

    doc.add_heading("3.3 Módulo 3: ApiGateway (Puerto 9056)", level=2)
    doc.add_paragraph(
        "Spring Cloud Gateway reactivo configurado con rutas dinámicas hacia el servicio descubierto por Eureka."
    )
    p_code3 = doc.add_paragraph()
    r = p_code3.add_run("Configuración application.yml:\n"
                      "server:\n"
                      "  port: 9056\n"
                      "spring:\n"
                      "  application:\n"
                      "    name: API-GATEWAY\n"
                      "  cloud:\n"
                      "    gateway:\n"
                      "      routes:\n"
                      "        - id: USER-SERVICE\n"
                      "          uri: lb://USER-SERVICE\n"
                      "          predicates:\n"
                      "            - Path=/client/**, /client\n"
                      "      default-filters:\n"
                      "        - DedupeResponseHeader=Access-Control-Allow-Credentials Access-Control-Allow-Origin\n"
                      "eureka:\n"
                      "  client:\n"
                      "    register-with-eureka: true\n"
                      "    fetch-registry: true\n"
                      "    service-url:\n"
                      "      defaultZone: http://localhost:9099/eureka")
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)

    doc.add_heading("4. Evidencias de Ejecución y Pruebas", level=1)

    table_tests = doc.add_table(rows=4, cols=4)
    table_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Prueba", "Endpoint / Recurso", "Resultado Esperado", "Estado"]
    for j, h in enumerate(headers):
        cell = table_tests.rows[0].cells[j]
        set_cell_background(cell, "2B6CB0")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    test_data = [
        ("1. Acceso Directo", "GET http://localhost:8086/client", "Welcome to client", "EXITOSO [OK]"),
        ("2. Enrutamiento Gateway", "GET http://localhost:9056/client", "Welcome to client", "EXITOSO [OK]"),
        ("3. Registro Eureka", "http://localhost:9099/eureka/apps", "USER-SERVICE y API-GATEWAY en estado UP", "EXITOSO [OK]")
    ]

    for i, row_data in enumerate(test_data):
        row = table_tests.rows[i+1]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if i % 2 == 1:
                set_cell_background(cell, "F7FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9.5)
            if j == 3:
                run.bold = True
                run.font.color.rgb = RGBColor(40, 167, 69)

    doc.add_paragraph()
    doc.add_heading("5. Conclusiones", level=1)
    doc.add_paragraph(
        "1. La implementación realizada cumple en su totalidad con el tutorial de GeeksforGeeks, manteniendo la misma estructura, nombres de paquetes, endpoints y propiedades de configuración.\n"
        "2. Se demostró el desacoplamiento arquitectónico: los clientes externos únicamente consumen el API Gateway en el puerto 9056, permitiendo escalar o reubicar microservicios sin impactar al cliente.\n"
        "3. El balanceo de carga automático provisto por Spring Cloud LoadBalancer junto con Eureka garantiza alta disponibilidad y resiliencia en la comunicación entre microservicios."
    )

    out_file = r"C:\Users\josed\taller-eureka-gateway\INFORME_ENTREGABLE_TALLER_EUREKA_GATEWAY.docx"
    doc.save(out_file)
    print(f"Report saved to {out_file}")

if __name__ == "__main__":
    create_report()
